import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from doc_design import (
    COLOR_NAVY, COLOR_SLATE, COLOR_TEXT, COLOR_MUTED, COLOR_ZEBRA,
    COLOR_CALLOUT_BG, COLOR_CALLOUT_BDR, COLOR_WARN_BG, COLOR_WARN_BDR,
    COLOR_CRIT_BG, COLOR_CRIT_BDR, COLOR_CODE_BG, COLOR_BORDER,
    set_cell_shading, set_cell_margins, add_header_block,
    h1, h2, h3, para, bullet, callout, code_box, table_builder
)

def add_callout_box(doc, text, title="NOTE", box_type="info"):
    callout(doc, text, title=title, box_type=box_type)

def add_code(doc, text):
    code_box(doc, text)

def add_table(doc, headers, rows, widths=None):
    return table_builder(doc, headers, rows, widths)

def add_image_if_exists(doc, path, caption=""):
    if os.path.exists(path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(path, width=Inches(6.4))
        
        if caption:
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(10)
            r_cap = p_cap.add_run(caption)
            r_cap.font.size = Pt(8)
            r_cap.font.italic = True
            r_cap.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        return True
    return False

def add_flowchart_card(doc, step_num, title, description, inputs="", outputs="", tools=""):
    """Render a visual pipeline card with high-contrast styling."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    set_cell_shading(c, "F8FAFC")
    set_cell_margins(c, top=80, bottom=80, left=100, right=100)
    
    tcPr = c._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{COLOR_SLATE}"/>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER}"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    r_step = p.add_run(f"STAGE {step_num}: ")
    r_step.font.bold = True
    r_step.font.size = Pt(9)
    r_step.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
    
    r_title = p.add_run(f"{title}\n")
    r_title.font.bold = True
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = RGBColor(0x0F, 0x29, 0x42)
    
    r_desc = p.add_run(f"{description}\n")
    r_desc.font.size = Pt(8.5)
    r_desc.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    
    meta_parts = []
    if inputs: meta_parts.append(f"Inputs: {inputs}")
    if tools: meta_parts.append(f"Engine: {tools}")
    if outputs: meta_parts.append(f"Artifacts: {outputs}")
    
    if meta_parts:
        r_meta = p.add_run(" | ".join(meta_parts))
        r_meta.font.size = Pt(7.5)
        r_meta.font.italic = True
        r_meta.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_flow_arrow(doc, label=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(f"▼  [{label}]  ▼" if label else "▼")
    r.font.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

print("Report generator base imported successfully.")
