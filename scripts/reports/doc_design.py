"""
doc_design.py
Shared styling, XML manipulation, and formatting helper functions for the
OneExtraction Master Executive Architecture & Project Report.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Executive Palette
COLOR_NAVY = "0F2942"       # Deep Executive Navy (#0F2942)
COLOR_SLATE = "2563EB"      # Tech Accent Slate/Royal Blue (#2563EB)
COLOR_TEXT = "1E293B"       # Dark Slate Body (#1E293B)
COLOR_MUTED = "64748B"      # Muted Slate (#64748B)
COLOR_ZEBRA = "F8FAFC"      # Light Table Zebra Tint (#F8FAFC)
COLOR_CALLOUT_BG = "F0F7FF" # Light Tint Blue (#F0F7FF)
COLOR_CALLOUT_BDR = "2563EB"# Vibrant Blue Border
COLOR_WARN_BG = "FFFBEB"    # Light Amber (#FFFBEB)
COLOR_WARN_BDR = "D97706"   # Amber (#D97706)
COLOR_CRIT_BG = "FEF2F2"    # Light Red (#FEF2F2)
COLOR_CRIT_BDR = "DC2626"   # Red (#DC2626)
COLOR_CODE_BG = "F1F5F9"    # Light Code Gray (#F1F5F9)
COLOR_BORDER = "CBD5E1"     # Slate Border (#CBD5E1)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'))

def set_cell_margins(cell, top=70, bottom=70, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_header_block(doc, title, subtitle, date_str="October 2026"):
    # Pre-title
    p_pre = doc.add_paragraph()
    p_pre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_pre = p_pre.add_run("ENTERPRISE SOLUTION ARCHITECTURE & TECHNICAL PROJECT REPORT")
    r_pre.font.size = Pt(9.5)
    r_pre.font.bold = True
    r_pre.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
    p_pre.paragraph_format.space_after = Pt(4)

    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.font.size = Pt(21)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x0F, 0x29, 0x42)
    p.paragraph_format.space_after = Pt(4)

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run(subtitle)
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    p_sub.paragraph_format.space_after = Pt(8)

    # Metadata Table
    tbl = doc.add_table(rows=6, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta = [
        ("Target Audience", "Senior Stakeholders, Business Leadership, Project Governance, Solution Architects, Tech Management"),
        ("Project Identifier", "OneExtraction Platform — B2B Lead Intelligence & OpenStreetMap Spatial Harvester"),
        ("Repository Location", "D:\\Data Scraping Project POC\\OneExtraction"),
        ("Production Version", "v2.1.0 Enterprise Production Ready (Docker / FastAPI / Next.js / PostgreSQL / PostGIS)"),
        ("Report Release Date", date_str),
        ("Engineering Status", "Active Production Baseline (Phase 1 Resilience & Dual-Pipeline Harvest Verified)"),
    ]
    for i, (k, v) in enumerate(meta):
        c0, c1 = tbl.rows[i].cells[0], tbl.rows[i].cells[1]
        set_cell_shading(c0, "F1F5F9")
        set_cell_shading(c1, "F8FAFC")
        set_cell_margins(c0, 45, 45, 80, 80)
        set_cell_margins(c1, 45, 45, 80, 80)
        c0.width = Inches(2.2)
        c1.width = Inches(4.5)
        p0 = c0.paragraphs[0]; p0.paragraph_format.space_after = Pt(1); p0.paragraph_format.space_before = Pt(1)
        r0 = p0.add_run(k); r0.font.bold = True; r0.font.size = Pt(8.5); r0.font.color.rgb = RGBColor(0x0F, 0x29, 0x42)
        p1 = c1.paragraphs[0]; p1.paragraph_format.space_after = Pt(1); p1.paragraph_format.space_before = Pt(1)
        r1 = p1.add_run(v); r1.font.size = Pt(8.5); r1.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    p_line = doc.add_paragraph()
    p_line.paragraph_format.space_before = Pt(10)
    p_line.paragraph_format.space_after = Pt(14)
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="{COLOR_NAVY}"/></w:pBdr>')
    p_line._p.get_or_add_pPr().append(pBdr)

def h1(doc, text):
    h = doc.add_heading(text, level=1)
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(4)
    for r in h.runs:
        r.font.size = Pt(12.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x0F, 0x29, 0x42)
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="{COLOR_SLATE}"/></w:pBdr>')
    h._p.get_or_add_pPr().append(pBdr)

def h2(doc, text):
    h = doc.add_heading(text, level=2)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(3)
    for r in h.runs:
        r.font.size = Pt(10.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

def h3(doc, text):
    h = doc.add_heading(text, level=3)
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(2)
    for r in h.runs:
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

def para(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    for r in p.runs:
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

def bullet(doc, text, bold_prefix=None, indent_level=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Inches(0.22 * (indent_level + 1))
    
    bullet_run = p.add_run("• ")
    bullet_run.font.bold = True
    bullet_run.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
    bullet_run.font.size = Pt(9)

    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.bold = True
        r_pre.font.size = Pt(9)
        r_pre.font.color.rgb = RGBColor(0x0F, 0x29, 0x42)
    
    r_body = p.add_run(text)
    r_body.font.size = Pt(9)
    r_body.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

def callout(doc, text, title="NOTE", box_type="info"):
    color = COLOR_CALLOUT_BG
    border_color = COLOR_CALLOUT_BDR
    text_color = RGBColor(0x0F, 0x29, 0x42)
    
    if box_type == "warning":
        color = COLOR_WARN_BG
        border_color = COLOR_WARN_BDR
        text_color = RGBColor(0xB4, 0x53, 0x09)
    elif box_type == "critical":
        color = COLOR_CRIT_BG
        border_color = COLOR_CRIT_BDR
        text_color = RGBColor(0xB9, 0x1C, 0x1C)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    set_cell_shading(c, color)
    set_cell_margins(c, top=90, bottom=90, left=120, right=120)
    
    tcPr = c._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    r_title = p.add_run(f"[{title}] ")
    r_title.font.bold = True
    r_title.font.size = Pt(8.5)
    r_title.font.color.rgb = text_color
    
    r_txt = p.add_run(text)
    r_txt.font.size = Pt(8.5)
    r_txt.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
    
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)

def code_box(doc, text_content):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.cell(0, 0)
    set_cell_shading(c, COLOR_CODE_BG)
    set_cell_margins(c, top=80, bottom=80, left=110, right=110)
    
    tcPr = c._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="12" w:space="0" w:color="{COLOR_BORDER}"/>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{COLOR_BORDER}"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="{COLOR_BORDER}"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="{COLOR_BORDER}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    r = p.add_run(text_content)
    r.font.name = "Consolas"
    r.font.size = Pt(7.5)
    r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(2)

def table_builder(doc, headers, data_rows, col_widths=None):
    tbl = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    h_row = tbl.rows[0]
    for idx, heading in enumerate(headers):
        cell = h_row.cells[idx]
        set_cell_shading(cell, COLOR_NAVY)
        set_cell_margins(cell, top=70, bottom=70, left=70, right=70)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(heading)
        r.font.bold = True
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for r_idx, row_values in enumerate(data_rows):
        row = tbl.rows[r_idx + 1]
        bg = COLOR_ZEBRA if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            cell = row.cells[c_idx]
            set_cell_shading(cell, bg)
            set_cell_margins(cell, top=55, bottom=55, left=65, right=65)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(val))
            r.font.size = Pt(7.2)
            r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    if col_widths:
        for row in tbl.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(3)
    return tbl

print("doc_design module loaded successfully.")
